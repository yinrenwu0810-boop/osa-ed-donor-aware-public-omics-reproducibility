from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "03_manuscript" / "manuscript_zh.v2.src.md"
OUTPUT = ROOT / "06_delivery" / "IH_OSA_ED_中文修订工作稿_v2.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

ACCENT = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side in ("top", "bottom", "start", "end"):
                tag = qn(f"w:{side}")
                node = margins.find(tag)
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), "80" if side in ("top", "bottom") else "120")
                node.set(qn("w:type"), "dxa")


def add_runs(paragraph, text, size=10.5, color=None):
    # Supports the manuscript's bold Markdown without changing citation syntax.
    cursor = 0
    for match in re.finditer(r"\*\*(.+?)\*\*", text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color)
        run = paragraph.add_run(match.group(1))
        set_run_font(run, size=size, bold=True, color=color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color)


def image_path(raw):
    raw = raw.replace("\\", "/")
    if raw.startswith("C:/Users/22394/Documents/Codex/2026-07-12/acad/"):
        return Path(raw.replace("C:/Users/22394/Documents/Codex/2026-07-12/acad/", str(ROOT.parent) + "/"))
    return (SOURCE.parent / raw).resolve()


def add_figure(doc, raw_path, alt):
    path = image_path(raw_path)
    if not path.is_file():
        note = doc.add_paragraph()
        note.paragraph_format.space_after = Pt(4)
        add_runs(note, f"[图件未嵌入：{alt}]", size=9.5, color=MUTED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.15))


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in [
        ("Heading 1", 16, ACCENT, 18, 10),
        ("Heading 2", 13, ACCENT, 12, 6),
        ("Heading 3", 12, DARK, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("IH–OSA–ED | v2 修订工作稿")
    set_run_font(run, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("v2 修订工作稿｜仅供内部修订与作者核对｜")
    set_run_font(run, size=8.5, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ")
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(88)
    cover.paragraph_format.space_after = Pt(12)
    run = cover.add_run(title)
    set_run_font(run, size=21, bold=True, color=DARK)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(32)
    add_runs(sub, "中文修订工作稿 v2｜基于 Gate 01–04 冻结结果", size=12, color=MUTED)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(8)
    add_runs(note, "状态说明：本文件供作者核对与后续投稿制品终审使用；尚未替代最终投稿稿。", size=10, color=RGBColor(122, 90, 0))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    i = 1
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            add_runs(p, line[2:], size=16, color=ACCENT)
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            add_runs(p, line[3:], size=13, color=ACCENT)
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_runs(p, line[4:], size=12, color=DARK)
        elif line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if match:
                add_figure(doc, match.group(2), match.group(1))
        elif line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header_cells = [x.strip() for x in line.strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header_cells))
            widths = [9360 // len(header_cells)] * len(header_cells)
            widths[-1] += 9360 - sum(widths)
            set_table_geometry(table, widths)
            for col, text in enumerate(header_cells):
                cell = table.cell(0, col)
                set_cell_shading(cell, "F4F6F9")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                add_runs(p, text, size=8.5, color=DARK)
                for run in p.runs:
                    run.bold = True
            for r_idx, row in enumerate(rows, start=1):
                for c_idx, text in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    add_runs(p, text, size=8.2)
            continue
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, line[2:], size=10.5)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.74)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.25
            add_runs(p, line, size=10.5)
        i += 1

    doc.core_properties.title = title
    doc.core_properties.subject = "中文修订工作稿 v2"
    doc.core_properties.author = ""
    doc.core_properties.comments = "Internal revision working draft; not submission-final."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_doc()
