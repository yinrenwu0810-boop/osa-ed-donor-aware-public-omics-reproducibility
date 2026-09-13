from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


HERE = Path(__file__).resolve().parent
REVISION = HERE.parents[1]
V3_SOURCE = HERE / "manuscript_zh.v3_virtual_perturbation.src.md"
V3_DOCX = REVISION / "06_delivery" / "IH_OSA_ED_中文修订工作稿_v3_含虚拟扰动结果.docx"
FIGURE = HERE / "figures" / "Figure5_virtual_perturbation_hypothesis_prioritization.png"
OUTPUT = REVISION / "06_delivery" / "IH_OSA_ED_中文修订工作稿_v4_含虚拟扰动结果与图5.docx"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    if OUTPUT.exists():
        raise SystemExit(f"Refusing to overwrite existing output: {OUTPUT}")
    for required in (V3_SOURCE, V3_DOCX, FIGURE):
        if not required.is_file():
            raise SystemExit(f"Missing required input: {required}")

    doc = Document(V3_DOCX)
    caption = next(
        (p for p in doc.paragraphs if p.text.startswith("图5｜虚拟扰动的跨供体稳定性")),
        None,
    )
    if caption is None:
        raise RuntimeError("Cannot locate the Figure 5 caption in the v3 manuscript")

    figure_paragraph = doc.add_paragraph()
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.paragraph_format.space_before = Pt(8)
    figure_paragraph.paragraph_format.space_after = Pt(4)
    figure_paragraph.paragraph_format.keep_with_next = True
    figure_paragraph.add_run().add_picture(str(FIGURE), width=Inches(5.50))

    caption.paragraph_format.keep_together = True
    caption.paragraph_format.keep_with_next = True
    caption._p.addprevious(figure_paragraph._p)

    for section in doc.sections:
        for paragraph in section.header.paragraphs + section.footer.paragraphs:
            for run in paragraph.runs:
                if "v3" in run.text:
                    run.text = run.text.replace("v3", "v4")
    doc.save(OUTPUT)

    print(f"v3_source_sha256={sha256(V3_SOURCE)}")
    print(f"v3_docx_sha256={sha256(V3_DOCX)}")
    print(f"figure_png_sha256={sha256(FIGURE)}")
    print(f"v4_docx_sha256={sha256(OUTPUT)}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
