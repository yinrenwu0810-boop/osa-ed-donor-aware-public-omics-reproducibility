#!/usr/bin/env python3
"""Build a new Chinese manuscript v3 from sealed v2 plus VP-G04--G07 insertions.

The v2 source and v2 DOCX are read only.  This script only creates versioned v3
artifacts after the figure has been generated.
"""

from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path

from docx import Document


ACADEMIC_ROOT = Path(__file__).resolve().parents[3]
REVISION = ACADEMIC_ROOT / "revision_v2"
HERE = Path(__file__).resolve().parent
V2_SOURCE = REVISION / "03_manuscript" / "manuscript_zh.v2.src.md"
INSERT = HERE / "VP_manuscript_insert_zh_v1.md"
V3_SOURCE = HERE / "manuscript_zh.v3_virtual_perturbation.src.md"
FIGURE = HERE / "figures" / "Figure5_virtual_perturbation_hypothesis_prioritization.png"
OUTPUT = REVISION / "06_delivery" / "IH_OSA_ED_中文修订工作稿_v3_含虚拟扰动结果.docx"
V2_OUTPUT = REVISION / "06_delivery" / "IH_OSA_ED_中文修订工作稿_v2.docx"
BASE_BUILDER = REVISION / "scripts" / "build_chinese_manuscript_v2_docx.py"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_source() -> None:
    base = V2_SOURCE.read_text(encoding="utf-8")
    insert = INSERT.read_text(encoding="utf-8").strip()
    if "## 2.8." in base or "## 4.10. 虚拟扰动" in base:
        raise RuntimeError("The sealed v2 source unexpectedly already contains the VP insertion")
    if "# 3. 讨论" not in base or "## 4.10. 软件与可重复性" not in base:
        raise RuntimeError("Expected v2 insertion anchors are absent")
    result_insert, method_insert, discussion_insert = insert.split("\n\n## 4.10. ", 1)[0], "## 4.10. " + insert.split("\n\n## 4.10. ", 1)[1].split("\n\n虚拟扰动分析将", 1)[0], "虚拟扰动分析将" + insert.split("\n\n虚拟扰动分析将", 1)[1]
    result_insert = result_insert.replace("\n\n虚拟扰动分析将候选", "\n\n虚拟扰动分析将候选")
    source = base.replace("# 3. 讨论", result_insert + "\n\n# 3. 讨论", 1)
    source = source.replace("## 4.10. 软件与可重复性", method_insert + "\n\n## 4.11. 软件与可重复性", 1)
    anchor = "后续验证可直接针对上述限制展开。"
    source = source.replace(anchor, discussion_insert + "\n\n" + anchor, 1)
    source = source.replace("../02_results/", "C:/Users/22394/Documents/Codex/2026-07-12/acad/revision_v2/02_results/")
    V3_SOURCE.write_text(source + "\n", encoding="utf-8")


def build_docx() -> None:
    spec = importlib.util.spec_from_file_location("sealed_v2_builder", BASE_BUILDER)
    if spec is None or spec.loader is None:
        raise RuntimeError("Cannot load the sealed v2 builder")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.SOURCE = V3_SOURCE
    module.OUTPUT = OUTPUT
    module.build_doc()
    doc = Document(OUTPUT)
    for section in doc.sections:
        for paragraph in section.header.paragraphs + section.footer.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace("v2 修订工作稿", "v3 修订工作稿（含虚拟扰动结果）")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace("中文修订工作稿 v2｜基于 Gate 01–04 冻结结果", "中文修订工作稿 v3｜含 VP-G00–G07 封存计算结果")
    doc.core_properties.subject = "中文修订工作稿 v3（含虚拟扰动结果）"
    doc.core_properties.comments = "Versioned update; v2 is retained unchanged. VP-G08 has no wet-lab data."
    doc.save(OUTPUT)


def main() -> None:
    if OUTPUT.exists() or V3_SOURCE.exists():
        raise SystemExit("Refusing to overwrite an existing v3 source or DOCX")
    if not (V2_SOURCE.is_file() and V2_OUTPUT.is_file() and INSERT.is_file() and FIGURE.is_file()):
        raise SystemExit("A required sealed input, insertion source, or figure is missing")
    build_source()
    build_docx()
    print(f"v2_source_sha256={sha256(V2_SOURCE)}")
    print(f"v2_docx_sha256={sha256(V2_OUTPUT)}")
    print(f"v3_source_sha256={sha256(V3_SOURCE)}")
    print(f"v3_docx_sha256={sha256(OUTPUT)}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
