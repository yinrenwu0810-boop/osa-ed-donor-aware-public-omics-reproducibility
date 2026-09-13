from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "VIRTUAL_PERTURBATION_WORKBOOK_zh.md"
OUTPUT = HERE / "TYMS_EFNB2_LRRC17_虚拟扰动实施工作手册_v1.docx"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
BORDER = "C9D2DE"
CAUTION = "FFF4CE"


def set_run_font(run, size=11, bold=False, italic=False, color="000000", mono=False):
    name = "Consolas" if mono else "Calibri"
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_table_row_no_split(row):
    """Keep a logical table row on one page whenever Word can do so."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_border_and_fill(paragraph, fill, color):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), color)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_custom_numbering(doc, num_format, text, left=540, hanging=270):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(existing_abs, default=-1) + 1
    num_id = max(existing_num, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    lvl.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abs_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_ref)
    p_pr.insert(0, num_pr)


def clean_md(text):
    text = re.sub(r"\[([^\]]+)\]\(<([^>]+)>\)", r"\1（\2）", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    text = re.sub(r"\[\^(\d+)\]", r"[\1]", text)
    return text


INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")


def add_inline(paragraph, text, size=11, color="000000"):
    text = clean_md(text)
    cursor = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=max(size - 0.5, 8), color=DARK_BLUE, mono=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True, color=color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color)


def choose_widths(headers):
    n = len(headers)
    joined = "|".join(headers)
    if n == 2:
        return [2700, 6660]
    if n == 3:
        return [1800, 2700, 4860]
    if n == 4:
        if "SHA-256" in joined:
            return [3200, 1300, 1800, 3060]
        return [1450, 1900, 2200, 3810]
    if n == 5:
        return [1050, 1350, 1350, 2250, 3360]
    return [CONTENT_WIDTH_DXA // n] * (n - 1) + [CONTENT_WIDTH_DXA - (CONTENT_WIDTH_DXA // n) * (n - 1)]


def add_table(doc, rows):
    headers = [x.strip() for x in rows[0]]
    widths = choose_widths(headers)
    table = doc.add_table(rows=len(rows), cols=len(headers))
    for i, row in enumerate(rows):
        set_table_row_no_split(table.rows[i])
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, value.strip(), size=9.2 if len(headers) >= 4 else 9.6)
            if i == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(INK)
        if i == 0:
            set_repeat_table_header(table.rows[i])
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("IH–OSA–ED  |  虚拟扰动实施工作手册")
    set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(p)


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("TYMS、EFNB2、LRRC17")
    set_run_font(run, size=18, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("单细胞虚拟扰动实施工作手册")
    set_run_font(run, size=30, bold=True, color=INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run("供新对话零上下文接手、逐 Gate 执行与审计")
    set_run_font(run, size=14, color=MUTED)
    for label, value in (
        ("项目", "IH–OSA–ED transcriptomic revision_v2"),
        ("版本", "1.0"),
        ("冻结日期", "2026-08-27"),
        ("执行边界", "虚拟扰动用于假设生成，不替代真实细胞因果验证"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{label}：")
        set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color="333333")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    r = p.add_run("权威源：VIRTUAL_PERTURBATION_WORKBOOK_zh.md")
    set_run_font(r, size=9.5, italic=True, color=MUTED)
    doc.add_page_break()


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)
    doc.core_properties.title = "TYMS、EFNB2、LRRC17 单细胞虚拟扰动实施工作手册"
    doc.core_properties.subject = "IH–OSA–ED revision_v2 virtual perturbation runbook"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "scTenifoldKnk, GEARS, UniPert, G2CP, TYMS, EFNB2, LRRC17"
    add_cover(doc)
    bullet_id = add_custom_numbering(doc, "bullet", "•")
    decimal_id = add_custom_numbering(doc, "decimal", "%1.")

    i = 0
    in_code = False
    code_lang = ""
    code_lines = []
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = line[3:].strip()
                code_lines = []
            else:
                if code_lang == "mermaid":
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(8)
                    set_paragraph_border_and_fill(p, LIGHT_BLUE, BLUE)
                    add_inline(p, "总流程：VP-G00 → VP-G01 → VP-G02 → VP-G03 → VP-G04 → VP-G05 → VP-G06 → VP-G07 → VP-G08", size=10, color=INK)
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.18)
                    p.paragraph_format.right_indent = Inches(0.08)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(7)
                    p.paragraph_format.line_spacing = 1.05
                    set_paragraph_border_and_fill(p, LIGHT_GRAY, BORDER)
                    run = p.add_run("\n".join(code_lines))
                    set_run_font(run, size=8.3, color="263238", mono=True)
                in_code = False
                code_lang = ""
                code_lines = []
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            parsed = []
            for idx, tline in enumerate(table_lines):
                cells = [x.strip() for x in tline.strip().strip("|").split("|")]
                if idx == 1 and all(re.fullmatch(r":?-{3,}:?", x.replace(" ", "")) for x in cells):
                    continue
                parsed.append(cells)
            if parsed:
                add_table(doc, parsed)
            continue
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, re.sub(r"^[^\w\u4e00-\u9fff]+\s*", "", line[3:]), size=16, color=BLUE)
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:], size=13, color=BLUE)
        elif line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, line[5:], size=12, color=DARK_BLUE)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.12)
            p.paragraph_format.right_indent = Inches(0.06)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(7)
            fill = CAUTION if "⚠️" in line or "警告" in line else CALLOUT
            color = "C58A00" if fill == CAUTION else BLUE
            set_paragraph_border_and_fill(p, fill, color)
            add_inline(p, line[2:], size=10.3, color=INK)
        elif re.match(r"^- \[[ xX]\] ", line):
            checked = line[3].lower() == "x"
            p = doc.add_paragraph()
            apply_num(p, bullet_id)
            add_inline(p, ("[已完成] " if checked else "[待检查] ") + line[6:], size=10.5)
            p.paragraph_format.space_after = Pt(4)
        elif line.startswith("- "):
            p = doc.add_paragraph()
            apply_num(p, bullet_id)
            add_inline(p, line[2:], size=10.5)
            p.paragraph_format.space_after = Pt(4)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph()
            apply_num(p, decimal_id)
            add_inline(p, re.sub(r"^\d+\. ", "", line), size=10.5)
            p.paragraph_format.space_after = Pt(4)
        elif re.match(r"^\[\^\d+\]:", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(5)
            add_inline(p, re.sub(r"^\[\^(\d+)\]:", r"[\1]", line), size=9.3, color=MUTED)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, line, size=10.8)
        i += 1

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
